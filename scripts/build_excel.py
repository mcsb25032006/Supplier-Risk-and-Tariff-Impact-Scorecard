"""
Excel Scorecard Workbook Builder
================================
Generates excel/Supplier_Risk_Tariff_Scorecard.xlsx using openpyxl.
Contains:
1. README: Project background, assumptions, and color legends
2. Scorecard: Live Excel formulas (INDEX/MATCH, SUMIFS, COUNTIFS, AVERAGEIFS, SUMPRODUCT population StDev, RAG formatting)
3. Tariff_Scenario_Toggle: Dynamic dropdown toggle calculating scenario landed spend
4. Data_Quality_Log: Real-time formulas auditing duplicates and missing costs
5. Raw_* Tabs: Source reference and transaction tables
"""

from pathlib import Path
from datetime import datetime
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.utils import get_column_letter

project_root = Path(__file__).resolve().parent.parent
data_dir = project_root / "data"
excel_dir = project_root / "excel"
excel_dir.mkdir(parents=True, exist_ok=True)

# ---------------------------------------------------------------------------
# Load source data
# ---------------------------------------------------------------------------
suppliers = pd.read_csv(data_dir / "suppliers.csv")
po = pd.read_csv(data_dir / "purchase_orders.csv")
tariff_ref = pd.read_csv(data_dir / "tariff_reference.csv")
tariff_scen = pd.read_csv(data_dir / "tariff_scenarios.csv")
defects = pd.read_csv(data_dir / "quality_defects.csv")


def to_date(s):
    if pd.isna(s) or s == "":
        return None
    return datetime.strptime(str(s), "%Y-%m-%d").date()


def build_excel_workbook():
    wb = Workbook()
    wb.remove(wb.active)

    FONT = "Arial"
    HEADER_FILL = PatternFill("solid", fgColor="1F3864")
    HEADER_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
    TITLE_FONT = Font(name=FONT, bold=True, size=14, color="1F3864")
    SUBTITLE_FONT = Font(name=FONT, italic=True, size=10, color="595959")
    BODY_FONT = Font(name=FONT, size=10)
    BOLD_FONT = Font(name=FONT, bold=True, size=10)
    THIN = Side(style="thin", color="D9D9D9")
    BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

    GREEN_FILL = PatternFill("solid", fgColor="C6EFCE")
    YELLOW_FILL = PatternFill("solid", fgColor="FFEB9C")
    RED_FILL = PatternFill("solid", fgColor="FFC7CE")
    GREEN_FONT = Font(name=FONT, color="006100", size=10)
    YELLOW_FONT = Font(name=FONT, color="9C6500", size=10)
    RED_FONT = Font(name=FONT, color="9C0006", size=10)
    INPUT_FILL = PatternFill("solid", fgColor="FFFF00")

    def style_header_row(ws, row, ncols):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=row, column=c)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = BORDER

    def autosize(ws, widths):
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w

    # ---------------------------------------------------------------------------
    # TAB 1: README
    # ---------------------------------------------------------------------------
    ws = wb.create_sheet("README")
    ws["A1"] = "Supplier Risk & Tariff-Impact Scorecard"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = "Datacenter hardware procurement scenario: GPU servers, network switches, PDU/MEP equipment"
    ws["A2"].font = SUBTITLE_FONT
    ws["A4"] = "How this workbook is organized"
    ws["A4"].font = BOLD_FONT
    notes = [
        ("Scorecard", "The primary deliverable: on-time rate, lead-time variability, landed cost, tariff exposure, defect rate, and a weighted composite risk score with RAG status, per supplier."),
        ("Tariff_Scenario_Toggle", "Pick a tariff scenario from the dropdown and see landed spend recalculate per supplier in real time."),
        ("Data_Quality_Log", "Live formulas auditing the raw export for duplicates and missing values prior to KPI calculation."),
        ("Raw_Suppliers", "Supplier master data (3 suppliers across US, Mexico, and China)."),
        ("Raw_POs", "344 purchase-order and shipment records simulating a Coupa export with intentional messiness for audit verification."),
        ("Raw_Tariff_Ref", "Current (July 2026) effective tariff rates by supplier country and HS code."),
        ("Raw_Tariff_Scenarios", "What-if scenarios feeding the dynamic dropdown on Tariff_Scenario_Toggle."),
        ("Raw_Defects", "Quality/defect log linked to POs flagged with defects."),
    ]
    r = 5
    for tab, desc in notes:
        ws.cell(row=r, column=1, value=tab).font = BOLD_FONT
        ws.cell(row=r, column=2, value=desc).font = BODY_FONT
        ws.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 28
        r += 1

    r += 1
    ws.cell(row=r, column=1, value="Color legend").font = BOLD_FONT
    r += 1
    legend = [
        (INPUT_FILL, "Yellow fill: input cell / scenario selector (edit this)"),
        (GREEN_FILL, "Green: Low risk (RAG)"),
        (YELLOW_FILL, "Yellow: Moderate risk (RAG)"),
        (RED_FILL, "Red: High risk (RAG)"),
    ]
    for fill, text in legend:
        ws.cell(row=r, column=1).fill = fill
        ws.cell(row=r, column=2, value=text).font = BODY_FONT
        r += 1

    r += 1
    ws.cell(row=r, column=1, value="Tariff assumptions (grounded in real July 2026 conditions)").font = BOLD_FONT
    r += 1
    assumptions = [
        "China: 35% effective rate = 10% Section 122 surcharge + 25% Section 301 (source: CBP/USTR guidance as tracked April 2026).",
        "Section 122's 10% surcharge is statutorily capped and sunsets ~Jul 24, 2026 unless renewed - modeled as a scenario.",
        "Mexico: USMCA-qualifying goods carry a ~0% effective rate; non-qualifying goods pay MFN rate.",
        "United States (domestic): 0% tariff exposure in all scenarios.",
    ]
    for a in assumptions:
        ws.cell(row=r, column=1, value="\u2022 " + a).font = BODY_FONT
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
        ws.row_dimensions[r].height = 15
        r += 1

    autosize(ws, [22, 90, 14, 14, 14, 14])

    # ---------------------------------------------------------------------------
    # TAB 2: Raw_Suppliers
    # ---------------------------------------------------------------------------
    ws = wb.create_sheet("Raw_Suppliers")
    headers = list(suppliers.columns)
    for c, h in enumerate(headers, start=1):
        ws.cell(row=1, column=c, value=h)
    style_header_row(ws, 1, len(headers))
    for i, row in suppliers.iterrows():
        for c, h in enumerate(headers, start=1):
            val = row[h]
            if h == "Onboarded_Date":
                val = to_date(val)
            cell = ws.cell(row=i + 2, column=c, value=val)
            cell.font = BODY_FONT
            if h == "Onboarded_Date":
                cell.number_format = "yyyy-mm-dd"
    autosize(ws, [12, 26, 15, 12, 16, 18, 16, 18])

    # ---------------------------------------------------------------------------
    # TAB 3: Raw_Tariff_Ref
    # ---------------------------------------------------------------------------
    ws = wb.create_sheet("Raw_Tariff_Ref")
    headers = list(tariff_ref.columns) + ["Key"]
    for c, h in enumerate(headers, start=1):
        ws.cell(row=1, column=c, value=h)
    style_header_row(ws, 1, len(headers))
    for i, row in tariff_ref.iterrows():
        rn = i + 2
        for c, h in enumerate(tariff_ref.columns, start=1):
            val = row[h]
            cell = ws.cell(row=rn, column=c, value=val)
            cell.font = BODY_FONT
            if h in ("Base_MFN_Rate", "Section_301_Rate", "Section_122_Surcharge", "Base_Effective_Rate"):
                cell.number_format = "0.0%"
        key_col = len(tariff_ref.columns) + 1
        ws.cell(row=rn, column=key_col, value=f'=A{rn}&"|"&B{rn}').font = BODY_FONT
    autosize(ws, [16, 12, 20, 15, 16, 20, 16, 16, 22])

    # ---------------------------------------------------------------------------
    # TAB 4: Raw_Tariff_Scenarios
    # ---------------------------------------------------------------------------
    ws = wb.create_sheet("Raw_Tariff_Scenarios")
    headers = list(tariff_scen.columns)
    for c, h in enumerate(headers, start=1):
        ws.cell(row=1, column=c, value=h)
    style_header_row(ws, 1, len(headers))
    for i, row in tariff_scen.iterrows():
        rn = i + 2
        for c, h in enumerate(headers, start=1):
            val = row[h]
            cell = ws.cell(row=rn, column=c, value=val)
            cell.font = BODY_FONT
            if h == "Adjustment_Value":
                cell.number_format = "0.0%;-0.0%"
    autosize(ws, [34, 16, 16, 55])

    # ---------------------------------------------------------------------------
    # TAB 5: Raw_Defects
    # ---------------------------------------------------------------------------
    ws = wb.create_sheet("Raw_Defects")
    headers = list(defects.columns)
    for c, h in enumerate(headers, start=1):
        ws.cell(row=1, column=c, value=h)
    style_header_row(ws, 1, len(headers))
    for i, row in defects.iterrows():
        rn = i + 2
        for c, h in enumerate(headers, start=1):
            val = row[h]
            if h == "Detected_Date":
                val = to_date(val)
            cell = ws.cell(row=rn, column=c, value=val)
            cell.font = BODY_FONT
            if h == "Detected_Date":
                cell.number_format = "yyyy-mm-dd"
    autosize(ws, [12, 12, 12, 22, 12, 15, 14])

    # ---------------------------------------------------------------------------
    # TAB 6: Raw_POs
    # ---------------------------------------------------------------------------
    ws = wb.create_sheet("Raw_POs")
    raw_cols = [
        "PO_ID", "Supplier_ID", "Product_Category", "HS_Code", "PO_Date", "Quantity",
        "Unit_Cost_USD", "Incoterm", "Freight_Mode", "Promised_Delivery_Date",
        "Actual_Delivery_Date", "Customs_Clearance_Days", "Defect_Flag", "PO_Status"
    ]
    helper_cols = [
        "Lead_Time_Days", "On_Time_Flag", "Supplier_Country", "Tariff_Key",
        "Effective_Tariff_Rate", "Base_Spend_USD", "Landed_Cost_Per_Unit", "Landed_Cost_Total",
        "Is_First_Occurrence"
    ]
    all_cols = raw_cols + helper_cols
    for c, h in enumerate(all_cols, start=1):
        ws.cell(row=1, column=c, value=h)
    style_header_row(ws, 1, len(all_cols))

    n = len(po)
    last_row = n + 1

    for i, row in po.iterrows():
        rn = i + 2
        ws.cell(row=rn, column=1, value=row["PO_ID"]).font = BODY_FONT
        ws.cell(row=rn, column=2, value=row["Supplier_ID"]).font = BODY_FONT
        ws.cell(row=rn, column=3, value=row["Product_Category"]).font = BODY_FONT
        ws.cell(row=rn, column=4, value=row["HS_Code"]).font = BODY_FONT
        d = ws.cell(row=rn, column=5, value=to_date(row["PO_Date"]))
        d.font = BODY_FONT
        d.number_format = "yyyy-mm-dd"
        ws.cell(row=rn, column=6, value=int(row["Quantity"])).font = BODY_FONT
        uc = row["Unit_Cost_USD"]
        cell_uc = ws.cell(row=rn, column=7, value=(None if pd.isna(uc) else float(uc)))
        cell_uc.font = BODY_FONT
        cell_uc.number_format = "$#,##0.00"
        ws.cell(row=rn, column=8, value=row["Incoterm"]).font = BODY_FONT
        ws.cell(row=rn, column=9, value=row["Freight_Mode"]).font = BODY_FONT
        d2 = ws.cell(row=rn, column=10, value=to_date(row["Promised_Delivery_Date"]))
        d2.font = BODY_FONT
        d2.number_format = "yyyy-mm-dd"
        actual_val = row["Actual_Delivery_Date"]
        d3 = ws.cell(row=rn, column=11, value=to_date(actual_val) if not pd.isna(actual_val) and actual_val != "" else None)
        d3.font = BODY_FONT
        d3.number_format = "yyyy-mm-dd"
        ws.cell(row=rn, column=12, value=int(row["Customs_Clearance_Days"])).font = BODY_FONT
        ws.cell(row=rn, column=13, value=row["Defect_Flag"]).font = BODY_FONT
        ws.cell(row=rn, column=14, value=row["PO_Status"]).font = BODY_FONT

        # Helper columns O-W
        ws.cell(row=rn, column=15, value=f'=IF(K{rn}="",0,K{rn}-E{rn})').font = BODY_FONT
        ws.cell(row=rn, column=16, value=f'=IF(N{rn}<>"Delivered","",IF(K{rn}<=J{rn},1,0))').font = BODY_FONT
        ws.cell(row=rn, column=17, value=f'=IFERROR(INDEX(Raw_Suppliers!$C$2:$C$4,MATCH(B{rn},Raw_Suppliers!$A$2:$A$4,0)),"")').font = BODY_FONT
        ws.cell(row=rn, column=18, value=f'=Q{rn}&"|"&D{rn}').font = BODY_FONT
        ws.cell(row=rn, column=19, value=f'=IFERROR(INDEX(Raw_Tariff_Ref!$H$2:$H$10,MATCH(R{rn},Raw_Tariff_Ref!$I$2:$I$10,0)),0)').font = BODY_FONT
        ws.cell(row=rn, column=19).number_format = "0.0%"
        ws.cell(row=rn, column=20, value=f'=IF(G{rn}="","",G{rn}*F{rn})').font = BODY_FONT
        ws.cell(row=rn, column=20).number_format = "$#,##0.00"
        ws.cell(row=rn, column=21, value=f'=IF(G{rn}="","",G{rn}*(1+S{rn}))').font = BODY_FONT
        ws.cell(row=rn, column=21).number_format = "$#,##0.00"
        ws.cell(row=rn, column=22, value=f'=IF(U{rn}="","",U{rn}*F{rn})').font = BODY_FONT
        ws.cell(row=rn, column=22).number_format = "$#,##0.00"
        ws.cell(row=rn, column=23, value=f'=IF(COUNTIF($A$2:A{rn},A{rn})=1,1,0)').font = BODY_FONT

    autosize(ws, [10, 11, 16, 10, 12, 10, 13, 10, 11, 14, 14, 11, 10, 11, 12, 11, 14, 16, 14, 14, 15, 15, 18])
    ws.freeze_panes = "A2"
    PO_LAST = last_row

    # ---------------------------------------------------------------------------
    # TAB 7: Data_Quality_Log
    # ---------------------------------------------------------------------------
    ws = wb.create_sheet("Data_Quality_Log")
    ws["A1"] = "Data Quality Audit: Raw_POs Export"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = "Live formulas auditing the raw Coupa export prior to KPI derivation."
    ws["A2"].font = SUBTITLE_FONT

    checks = [
        ("Total PO rows in export", f"=COUNTA(Raw_POs!A2:A{PO_LAST})"),
        ("Unique PO_IDs", f"=SUMPRODUCT(1/COUNTIF(Raw_POs!A2:A{PO_LAST},Raw_POs!A2:A{PO_LAST}))"),
        ("Duplicate row count (rows minus unique IDs)", "=B4-B5"),
        ("Rows missing Unit_Cost_USD (pending invoice reconciliation)", f"=COUNTBLANK(Raw_POs!G2:G{PO_LAST})"),
        ("POs still In Transit (excluded from on-time/lead-time KPIs)", f'=COUNTIF(Raw_POs!N2:N{PO_LAST},"In Transit")'),
        ("Inconsistent-case Incoterm values found", f'=SUMPRODUCT(--(EXACT(Raw_POs!H2:H{PO_LAST},UPPER(Raw_POs!H2:H{PO_LAST}))=FALSE))'),
    ]
    r = 4
    ws.cell(row=3, column=1, value="Audit Check").font = BOLD_FONT
    ws.cell(row=3, column=2, value="Result").font = BOLD_FONT
    style_header_row(ws, 3, 2)
    for label, formula in checks:
        ws.cell(row=r, column=1, value=label).font = BODY_FONT
        ws.cell(row=r, column=2, value=formula).font = BODY_FONT
        r += 1

    ws["A11"] = "Data Cleaning & Deduplication Architecture"
    ws["A11"].font = BOLD_FONT
    ws["A12"] = (
        "The raw export simulates typical ERP portal noise (re-keyed orders and unreconciled invoice lines). "
        "Raw_POs column W (Is_First_Occurrence) flags duplicate entries to ensure every Scorecard formula, SQL query, "
        "and Python pipeline step computes strictly identical KPIs across all deliverables."
    )
    ws["A12"].font = BODY_FONT
    ws["A12"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("A12:F12")
    ws.row_dimensions[12].height = 40
    autosize(ws, [55, 14, 12, 12, 12, 12])

    # ---------------------------------------------------------------------------
    # TAB 8: Scorecard
    # ---------------------------------------------------------------------------
    ws = wb.create_sheet("Scorecard")
    ws["A1"] = "Supplier Risk & Tariff-Impact Scorecard"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = "Weighted composite risk score: 30% delivery reliability, 20% lead-time volatility, 25% tariff exposure, 25% quality"
    ws["A2"].font = SUBTITLE_FONT

    sc_headers = [
        "Supplier_ID", "Supplier_Name", "Country", "Delivered_POs", "On_Time_POs",
        "On_Time_Rate", "Avg_Lead_Time_Days", "Lead_Time_StDev_Days", "Total_Base_Spend",
        "Total_Landed_Spend", "Tariff_Exposure_Pct", "Defect_Rate_Pct", "Avg_Customs_Days",
        "Composite_Risk_Score", "RAG_Status"
    ]
    header_row = 4
    for c, h in enumerate(sc_headers, start=1):
        ws.cell(row=header_row, column=c, value=h.replace("_", " "))
    style_header_row(ws, header_row, len(sc_headers))

    sup_ids = suppliers["Supplier_ID"].tolist()
    first_sc_row = header_row + 1

    for i, sid in enumerate(sup_ids):
        rn = first_sc_row + i
        ws.cell(row=rn, column=1, value=sid).font = BODY_FONT
        ws.cell(row=rn, column=2, value=f'=INDEX(Raw_Suppliers!$B$2:$B$4,MATCH(A{rn},Raw_Suppliers!$A$2:$A$4,0))').font = BODY_FONT
        ws.cell(row=rn, column=3, value=f'=INDEX(Raw_Suppliers!$C$2:$C$4,MATCH(A{rn},Raw_Suppliers!$A$2:$A$4,0))').font = BODY_FONT
        ws.cell(row=rn, column=4, value=(f'=COUNTIFS(Raw_POs!$B$2:$B${PO_LAST},A{rn},Raw_POs!$N$2:$N${PO_LAST},"Delivered",Raw_POs!$W$2:$W${PO_LAST},1)')).font = BODY_FONT
        ws.cell(row=rn, column=5, value=(f'=SUMIFS(Raw_POs!$P$2:$P${PO_LAST},Raw_POs!$B$2:$B${PO_LAST},A{rn},Raw_POs!$W$2:$W${PO_LAST},1)')).font = BODY_FONT
        f_cell = ws.cell(row=rn, column=6, value=f'=IFERROR(E{rn}/D{rn},0)')
        f_cell.font = BODY_FONT
        f_cell.number_format = "0.0%"
        g_cell = ws.cell(row=rn, column=7, value=(f'=IFERROR(AVERAGEIFS(Raw_POs!$O$2:$O${PO_LAST},Raw_POs!$B$2:$B${PO_LAST},A{rn},Raw_POs!$N$2:$N${PO_LAST},"Delivered",Raw_POs!$W$2:$W${PO_LAST},1),0)'))
        g_cell.font = BODY_FONT
        g_cell.number_format = "0.0"
        h_formula = (f'=IFERROR(SQRT(SUMPRODUCT((Raw_POs!$B$2:$B${PO_LAST}=A{rn})*(Raw_POs!$N$2:$N${PO_LAST}="Delivered")*(Raw_POs!$W$2:$W${PO_LAST}=1)*(Raw_POs!$O$2:$O${PO_LAST})^2)/D{rn}-G{rn}^2),0)')
        h_cell = ws.cell(row=rn, column=8, value=h_formula)
        h_cell.font = BODY_FONT
        h_cell.number_format = "0.0"
        i_cell = ws.cell(row=rn, column=9, value=(f'=IFERROR(SUMIFS(Raw_POs!$T$2:$T${PO_LAST},Raw_POs!$B$2:$B${PO_LAST},A{rn},Raw_POs!$W$2:$W${PO_LAST},1),0)'))
        i_cell.font = BODY_FONT
        i_cell.number_format = "$#,##0"
        j_cell = ws.cell(row=rn, column=10, value=(f'=IFERROR(SUMIFS(Raw_POs!$V$2:$V${PO_LAST},Raw_POs!$B$2:$B${PO_LAST},A{rn},Raw_POs!$W$2:$W${PO_LAST},1),0)'))
        j_cell.font = BODY_FONT
        j_cell.number_format = "$#,##0"
        k_cell = ws.cell(row=rn, column=11, value=f'=IFERROR((J{rn}-I{rn})/I{rn},0)')
        k_cell.font = BODY_FONT
        k_cell.number_format = "0.0%"
        l_cell = ws.cell(row=rn, column=12, value=(f'=IFERROR(COUNTIFS(Raw_POs!$B$2:$B${PO_LAST},A{rn},Raw_POs!$N$2:$N${PO_LAST},"Delivered",Raw_POs!$M$2:$M${PO_LAST},"Y",Raw_POs!$W$2:$W${PO_LAST},1)/D{rn},0)'))
        l_cell.font = BODY_FONT
        l_cell.number_format = "0.0%"
        m_cell = ws.cell(row=rn, column=13, value=(f'=IFERROR(AVERAGEIFS(Raw_POs!$L$2:$L${PO_LAST},Raw_POs!$B$2:$B${PO_LAST},A{rn},Raw_POs!$W$2:$W${PO_LAST},1),0)'))
        m_cell.font = BODY_FONT
        m_cell.number_format = "0.0"
        n_formula = f'=ROUND(0.30*(1-F{rn})*100+0.20*MIN(H{rn}/20*100,100)+0.25*K{rn}*100+0.25*L{rn}*100,1)'
        n_cell = ws.cell(row=rn, column=14, value=n_formula)
        n_cell.font = BODY_FONT
        n_cell.number_format = "0.0"
        o_cell = ws.cell(row=rn, column=15, value=f'=IF(N{rn}>=30,"Red",IF(N{rn}>=15,"Yellow","Green"))')
        o_cell.font = BODY_FONT
        o_cell.alignment = Alignment(horizontal="center")

    last_sc_row = first_sc_row + len(sup_ids) - 1

    # Totals row
    tot_row = last_sc_row + 1
    ws.cell(row=tot_row, column=2, value="TOTAL / BLENDED").font = BOLD_FONT
    ws.cell(row=tot_row, column=4, value=f"=SUM(D{first_sc_row}:D{last_sc_row})").font = BOLD_FONT
    ws.cell(row=tot_row, column=5, value=f"=SUM(E{first_sc_row}:E{last_sc_row})").font = BOLD_FONT
    tf = ws.cell(row=tot_row, column=6, value=f"=IFERROR(E{tot_row}/D{tot_row},0)")
    tf.font = BOLD_FONT
    tf.number_format = "0.0%"
    ti = ws.cell(row=tot_row, column=9, value=f"=SUM(I{first_sc_row}:I{last_sc_row})")
    ti.font = BOLD_FONT
    ti.number_format = "$#,##0"
    tj = ws.cell(row=tot_row, column=10, value=f"=SUM(J{first_sc_row}:J{last_sc_row})")
    tj.font = BOLD_FONT
    tj.number_format = "$#,##0"
    tk = ws.cell(row=tot_row, column=11, value=f"=IFERROR((J{tot_row}-I{tot_row})/I{tot_row},0)")
    tk.font = BOLD_FONT
    tk.number_format = "0.0%"

    # Conditional formatting
    ws.conditional_formatting.add(
        f"N{first_sc_row}:N{last_sc_row}",
        ColorScaleRule(start_type="min", start_color="C6EFCE", mid_type="percentile", mid_value=50, mid_color="FFEB9C", end_type="max", end_color="FFC7CE")
    )
    for rng in [f"O{first_sc_row}:O{last_sc_row}"]:
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Red"'], fill=RED_FILL, font=RED_FONT))
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Yellow"'], fill=YELLOW_FILL, font=YELLOW_FONT))
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Green"'], fill=GREEN_FILL, font=GREEN_FONT))

    for rn in range(first_sc_row, last_sc_row + 1):
        for c in range(1, len(sc_headers) + 1):
            ws.cell(row=rn, column=c).border = BORDER

    autosize(ws, [12, 24, 15, 13, 12, 12, 16, 17, 15, 15, 15, 13, 14, 16, 11])
    ws.freeze_panes = f"A{first_sc_row}"

    SCORECARD_FIRST = first_sc_row

    # ---------------------------------------------------------------------------
    # TAB 9: Tariff_Scenario_Toggle
    # ---------------------------------------------------------------------------
    ws = wb.create_sheet("Tariff_Scenario_Toggle")
    ws["A1"] = "Tariff Scenario Simulation Toggle"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = "Select a scenario from the dropdown to recalculate landed spend across all suppliers."
    ws["A2"].font = SUBTITLE_FONT

    ws["A4"] = "Select Tariff Scenario:"
    ws["A4"].font = BOLD_FONT
    ws["B4"] = "Current (Jul 2026)"
    ws["B4"].fill = INPUT_FILL
    ws["B4"].font = BOLD_FONT
    ws["B4"].border = BORDER

    dv = DataValidation(type="list", formula1="=Raw_Tariff_Scenarios!$A$2:$A$6", allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(ws["B4"])

    ws["D4"] = "Adjustment Type:"
    ws["D4"].font = BODY_FONT
    ws["E4"] = '=IFERROR(INDEX(Raw_Tariff_Scenarios!$B$2:$B$6,MATCH($B$4,Raw_Tariff_Scenarios!$A$2:$A$6,0)),"")'
    ws["E4"].font = BODY_FONT
    ws["F4"] = "Adjustment Value:"
    ws["F4"].font = BODY_FONT
    ws["G4"] = '=IFERROR(INDEX(Raw_Tariff_Scenarios!$C$2:$C$6,MATCH($B$4,Raw_Tariff_Scenarios!$A$2:$A$6,0)),0)'
    ws["G4"].font = BODY_FONT
    ws["G4"].number_format = "0.0%;-0.0%"
    ws["H4"] = "Scenario Description:"
    ws["H4"].font = BODY_FONT
    ws["I4"] = '=IFERROR(INDEX(Raw_Tariff_Scenarios!$D$2:$D$6,MATCH($B$4,Raw_Tariff_Scenarios!$A$2:$A$6,0)),"")'
    ws["I4"].font = BODY_FONT
    ws.merge_cells("I4:N4")

    tt_headers = [
        "Supplier_ID", "Supplier_Name", "Country", "Base_Spend", "Current_Effective_Rate",
        "Scenario_New_Rate", "Scenario_Landed_Spend", "Delta_vs_Current_Landed_Spend"
    ]
    header_row = 6
    for c, h in enumerate(tt_headers, start=1):
        ws.cell(row=header_row, column=c, value=h.replace("_", " "))
    style_header_row(ws, header_row, len(tt_headers))

    first_tt_row = header_row + 1
    for i in range(len(sup_ids)):
        rn = first_tt_row + i
        sc_row = SCORECARD_FIRST + i
        ws.cell(row=rn, column=1, value=f"=Scorecard!A{sc_row}").font = BODY_FONT
        ws.cell(row=rn, column=2, value=f"=Scorecard!B{sc_row}").font = BODY_FONT
        ws.cell(row=rn, column=3, value=f"=Scorecard!C{sc_row}").font = BODY_FONT
        d_cell = ws.cell(row=rn, column=4, value=f"=Scorecard!I{sc_row}")
        d_cell.font = BODY_FONT
        d_cell.number_format = "$#,##0"
        e_cell = ws.cell(row=rn, column=5, value=f"=Scorecard!K{sc_row}")
        e_cell.font = BODY_FONT
        e_cell.number_format = "0.0%"
        f_formula = f'=IF(C{rn}="United States",0,IF($E$4="Additive",MAX(0,E{rn}+$G$4),$G$4))'
        f_cell = ws.cell(row=rn, column=6, value=f_formula)
        f_cell.font = BODY_FONT
        f_cell.number_format = "0.0%"
        g_cell = ws.cell(row=rn, column=7, value=f"=D{rn}*(1+F{rn})")
        g_cell.font = BODY_FONT
        g_cell.number_format = "$#,##0"
        h_cell = ws.cell(row=rn, column=8, value=f"=G{rn}-D{rn}*(1+E{rn})")
        h_cell.font = BODY_FONT
        h_cell.number_format = "$#,##0;($#,##0)"

    last_tt_row = first_tt_row + len(sup_ids) - 1
    tot_row = last_tt_row + 1
    ws.cell(row=tot_row, column=2, value="TOTAL").font = BOLD_FONT
    td = ws.cell(row=tot_row, column=4, value=f"=SUM(D{first_tt_row}:D{last_tt_row})")
    td.font = BOLD_FONT
    td.number_format = "$#,##0"
    tg = ws.cell(row=tot_row, column=7, value=f"=SUM(G{first_tt_row}:G{last_tt_row})")
    tg.font = BOLD_FONT
    tg.number_format = "$#,##0"
    th = ws.cell(row=tot_row, column=8, value=f"=SUM(H{first_tt_row}:H{last_tt_row})")
    th.font = BOLD_FONT
    th.number_format = "$#,##0;($#,##0)"

    ws.conditional_formatting.add(
        f"H{first_tt_row}:H{tot_row}",
        ColorScaleRule(start_type="min", start_color="C6EFCE", mid_type="num", mid_value=0, mid_color="FFEB9C", end_type="max", end_color="FFC7CE")
    )

    for rn in range(first_tt_row, tot_row + 1):
        for c in range(1, len(tt_headers) + 1):
            ws.cell(row=rn, column=c).border = BORDER

    autosize(ws, [12, 24, 15, 14, 18, 16, 18, 22])
    ws.column_dimensions["I"].width = 40

    # ---------------------------------------------------------------------------
    # Sheet order + save
    # ---------------------------------------------------------------------------
    order = ["README", "Scorecard", "Tariff_Scenario_Toggle", "Data_Quality_Log",
             "Raw_Suppliers", "Raw_POs", "Raw_Tariff_Ref", "Raw_Tariff_Scenarios", "Raw_Defects"]
    wb._sheets = [wb[name] for name in order]
    for name in order:
        wb[name].sheet_view.showGridLines = False if name in ("README",) else True

    output_file = excel_dir / "Supplier_Risk_Tariff_Scorecard.xlsx"
    wb.save(output_file)
    print(f"Workbook successfully saved at: {output_file}")


if __name__ == "__main__":
    build_excel_workbook()
